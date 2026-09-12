import copy
from io import BytesIO
from uuid import UUID, uuid4

from docx import Document
from solution_copilot.application import deliveries
from solution_copilot.application.delivery_schemas import Coverage, QualityIssue, QualityOutput
from solution_copilot.domain.models import (
    CustomerAccess,
    DocumentChunk,
    GenerationRun,
    RequirementProfile,
    Solution,
)
from solution_copilot.infrastructure import documents as storage
from solution_copilot.infrastructure import exports, generation
from sqlalchemy import delete, select
from test_s01 import db as db
from test_s01 import setup as setup
from test_s07 import ready_solution


def completed_solution(setup, db, monkeypatch):
    client, h, customers, users, orgs, doc, solution, _ = ready_solution(setup, db, monkeypatch)
    row = db.get(Solution, UUID(solution["id"]))
    version = deliveries.solutions.current(db, row)
    chunk = db.scalar(select(DocumentChunk).where(DocumentChunk.document_id == doc.id))

    rich_content = {
        "type": "doc",
        "content": [
            {
                "type": "heading",
                "attrs": {"level": 2},
                "content": [{"type": "text", "text": "交付细节"}],
            },
            {
                "type": "paragraph",
                "content": [
                    {"type": "text", "text": "已确认正文 "},
                    {"type": "text", "text": "加粗", "marks": [{"type": "bold"}]},
                    {"type": "text", "text": "斜体", "marks": [{"type": "italic"}]},
                    {"type": "text", "text": "删除", "marks": [{"type": "strike"}]},
                    {"type": "text", "text": "代码", "marks": [{"type": "code"}]},
                    {"type": "text", "text": " <script>alert(1)</script> [1]"},
                ],
            },
            {
                "type": "bulletList",
                "content": [
                    {
                        "type": "listItem",
                        "content": [
                            {
                                "type": "paragraph",
                                "content": [{"type": "text", "text": "一级列表"}],
                            },
                            {
                                "type": "orderedList",
                                "attrs": {"start": 2, "type": "1"},
                                "content": [
                                    {
                                        "type": "listItem",
                                        "content": [
                                            {
                                                "type": "paragraph",
                                                "content": [{"type": "text", "text": "二级列表"}],
                                            }
                                        ],
                                    }
                                ],
                            },
                        ],
                    },
                    {
                        "type": "listItem",
                        "content": [{"type": "paragraph"}],
                    },
                ],
            },
            {
                "type": "codeBlock",
                "attrs": {"language": "python"},
                "content": [{"type": "text", "text": "print('<safe>')"}],
            },
            {
                "type": "table",
                "content": [
                    {
                        "type": "tableRow",
                        "content": [
                            {
                                "type": "tableHeader",
                                "attrs": {
                                    "colspan": 1,
                                    "rowspan": 1,
                                    "colwidth": None,
                                    "align": None,
                                },
                                "content": [
                                    {
                                        "type": "paragraph",
                                        "content": [
                                            {
                                                "type": "text",
                                                "text": "控制项",
                                                "marks": [{"type": "bold"}],
                                            }
                                        ],
                                    }
                                ],
                            },
                            {
                                "type": "tableHeader",
                                "content": [
                                    {
                                        "type": "paragraph",
                                        "content": [{"type": "text", "text": "状态"}],
                                    }
                                ],
                            },
                        ],
                    },
                    {
                        "type": "tableRow",
                        "content": [
                            {
                                "type": "tableCell",
                                "content": [
                                    {
                                        "type": "paragraph",
                                        "content": [
                                            {"type": "text", "text": "单点登录"},
                                            {"type": "hardBreak"},
                                            {"type": "text", "text": "权限隔离"},
                                        ],
                                    }
                                ],
                            },
                            {
                                "type": "tableCell",
                                "content": [
                                    {
                                        "type": "paragraph",
                                        "content": [
                                            {
                                                "type": "text",
                                                "text": "待核实",
                                                "marks": [{"type": "italic"}],
                                            }
                                        ],
                                    }
                                ],
                            },
                        ],
                    },
                ],
            },
        ],
    }
    version.sections = [
        {
            **section,
            "status": "draft",
            "content": rich_content
            if index == 0
            else {
                "type": "doc",
                "content": [
                    {
                        "type": "paragraph",
                        "content": [{"type": "text", "text": f"已确认正文 {section['title']} [1]"}],
                    }
                ],
            },
            "citations": [
                {
                    "id": "1",
                    "claim_text": "已确认正文",
                    "quote": "产品支持单点登录",
                    "source_type": "document",
                    "source_id": str(chunk.id),
                    "title": doc.title,
                    "document_id": str(doc.id),
                    "chunk_id": str(chunk.id),
                    "locator": {
                        "page_number": 1,
                        "section_path": ["安全"],
                        "generation": doc.generation,
                    },
                    "verification_status": "partial",
                }
            ],
            "warnings": [],
        }
        for index, section in enumerate(version.sections)
    ]
    row.status = "draft"
    db.commit()
    return client, h, customers, users, orgs, doc, solution


def start(client, h, solution, kind, who="a-owner", request_id=None, **extra):
    body = {
        "version": 1,
        "request_id": request_id or str(uuid4()),
        **extra,
    }
    response = client.post(f"/api/v1/solutions/{solution['id']}/{kind}", headers=h(who), json=body)
    return response, body


def set_missing_reference(db, solution, enabled):
    row = db.get(Solution, UUID(solution["id"]))
    version = deliveries.solutions.current(db, row)
    sections = copy.deepcopy(version.sections)
    text = sections[0]["content"]["content"][0]["content"][0]["text"]
    sections[0]["content"]["content"][0]["content"][0]["text"] = (
        text + " [9]" if enabled else text.removesuffix(" [9]")
    )
    version.sections = sections
    db.commit()


def test_verify_validates_locations_coverage_and_request_identity(setup, db, monkeypatch):
    client, h, _, _, _, _, solution = completed_solution(setup, db, monkeypatch)
    profile = db.scalar(
        select(RequirementProfile).where(
            RequirementProfile.project_id == UUID(solution["project_id"])
        )
    )
    section_id = solution["current"]["sections"][0]["id"]

    bad, _ = start(client, h, solution, "verify")
    monkeypatch.setattr(
        generation,
        "verify",
        lambda _: (
            QualityOutput(
                issues=[
                    QualityIssue(
                        rule_id="contradiction",
                        severity="error",
                        section_id="missing",
                        explanation="定位不存在",
                        suggestion="重试",
                    )
                ],
                coverage=[Coverage(requirement_id=i["id"], section_ids=[]) for i in profile.items],
            ),
            {"model": "test"},
        ),
    )
    db.rollback()
    deliveries.run_job(bad.json()["run"]["id"])
    db.expire_all()
    failed = db.get(GenerationRun, UUID(bad.json()["run"]["id"]))
    assert failed.status == "failed" and failed.error_message.startswith("模型校验定位")
    assert client.get(f"/api/v1/exports/{failed.id}/download", headers=h()).status_code == 404

    set_missing_reference(db, solution, True)
    output = QualityOutput(
        issues=[
            QualityIssue(
                rule_id="acceptance_metric",
                severity="warning",
                section_id=section_id,
                explanation="验收指标不够具体",
                suggestion="补充量化阈值",
            )
        ],
        coverage=[
            Coverage(requirement_id=i["id"], section_ids=[section_id]) for i in profile.items
        ],
    )
    monkeypatch.setattr(generation, "verify", lambda _: (output, {"model": "test"}))
    request_id = str(uuid4())
    queued, body = start(client, h, solution, "verify", request_id=request_id)
    assert queued.status_code == 202
    repeated, _ = start(client, h, solution, "verify", request_id=request_id)
    assert repeated.json()["run"]["id"] == queued.json()["run"]["id"]
    collision, _ = start(client, h, solution, "exports", request_id=request_id, format="markdown")
    assert collision.status_code == 409
    db.rollback()
    deliveries.run_job(queued.json()["run"]["id"])

    listed = client.get(
        f"/api/v1/solutions/{solution['id']}/deliveries?version=1", headers=h()
    ).json()
    result = next(item for item in listed if item["run"]["id"] == queued.json()["run"]["id"])
    assert result["coverage_percent"] == 100.0
    assert result["profile_version"] == profile.version
    assert result["issues"][0]["section_id"] == section_id
    assert any(
        issue["rule_id"] == "factual_citation" and issue["severity"] == "error"
        for issue in result["issues"]
    )
    assert {item["requirement_id"] for item in result["coverage"]} == {
        item["id"] for item in profile.items
    }
    assert body["version"] == result["version"]


def test_exports_are_openable_immutable_and_private(setup, db, monkeypatch):
    client, h, _, users, _, _, solution = completed_solution(setup, db, monkeypatch)
    objects = {}
    monkeypatch.setattr(
        storage, "put_original", lambda key, data, mime: objects.update({key: data})
    )
    monkeypatch.setattr(storage, "read_original", lambda key: objects[key])
    monkeypatch.setattr(storage, "delete_original", lambda key: objects.pop(key, None))

    markdown, _ = start(client, h, solution, "exports", format="markdown")
    section_id = solution["current"]["sections"][0]["id"]
    edited = client.patch(
        f"/api/v1/solutions/{solution['id']}",
        headers=h(),
        json={
            "version": 1,
            "request_id": str(uuid4()),
            "section_id": section_id,
            "content": {
                "type": "doc",
                "content": [
                    {"type": "paragraph", "content": [{"type": "text", "text": "新版本正文"}]}
                ],
            },
            "change_summary": "导出排队后编辑",
        },
    )
    assert edited.status_code == 200
    db.rollback()
    deliveries.run_job(markdown.json()["run"]["id"])
    url = markdown.json()["run"]["id"]
    owner_download = client.get(f"/api/v1/exports/{url}/download", headers=h())
    assert owner_download.status_code == 200
    text = owner_download.content.decode()
    assert "已确认正文" in text and "新版本正文" not in text
    assert "参考资料" in text and "产品支持单点登录" in text
    assert "### 交付细节" in text and "**加粗**" in text and "*斜体*" in text
    assert "- 一级列表" in text and "2. 二级列表" in text
    assert "\n- \n" in text
    assert "```python\nprint('<safe>')\n```" in text
    assert "| **控制项** | 状态 |" in text and "单点登录" in text and "权限隔离" in text
    assert "<script>" not in text and r"\<script\>alert\(1\)\</script\>" in text
    assert owner_download.headers["cache-control"] == "private, no-store"
    assert (
        client.get(f"/api/v1/exports/{url}/download", headers=h("a-viewer")).content
        == owner_download.content
    )
    assert client.get(f"/api/v1/exports/{url}/download", headers=h("b-owner", 1)).status_code == 404
    assert (
        start(client, h, solution, "exports", who="a-viewer", format="docx")[0].status_code == 403
    )

    docx, _ = start(client, h, solution, "exports", who="a-member", format="docx")
    db.rollback()
    deliveries.run_job(docx.json()["run"]["id"])
    docx_id = docx.json()["run"]["id"]
    downloaded = client.get(f"/api/v1/exports/{docx_id}/download", headers=h("a-member"))
    opened = Document(BytesIO(downloaded.content))
    assert downloaded.status_code == 200 and opened.paragraphs[0].text == "身份治理方案"
    assert any("产品支持单点登录" in p.text for p in opened.paragraphs)
    assert any(p.style.name == "Heading 2" and p.text == "交付细节" for p in opened.paragraphs)
    assert any(
        p.style.name == "No Spacing" and "print('<safe>')" in p.text for p in opened.paragraphs
    )
    assert any("• 一级列表" in p.text for p in opened.paragraphs)
    assert any("2. 二级列表" in p.text for p in opened.paragraphs)
    table = opened.tables[0]
    assert [[cell.text for cell in row.cells] for row in table.rows] == [
        ["控制项", "状态"],
        ["单点登录\n权限隔离", "待核实"],
    ]
    assert table.cell(0, 0).paragraphs[0].runs[0].bold
    assert table.cell(1, 1).paragraphs[0].runs[0].italic
    db.execute(delete(CustomerAccess).where(CustomerAccess.user_id == users["a-member"][0].id))
    db.commit()
    assert (
        client.get(f"/api/v1/exports/{docx_id}/download", headers=h("a-member")).status_code == 404
    )


def test_failed_cancelled_and_late_exports_never_download(setup, db, monkeypatch):
    client, h, _, _, _, _, solution = completed_solution(setup, db, monkeypatch)
    set_missing_reference(db, solution, True)
    missing, _ = start(client, h, solution, "exports", format="markdown")
    assert missing.status_code == 409 and missing.json()["error"]["code"] == "CITATION_MISSING"
    set_missing_reference(db, solution, False)

    failed, _ = start(client, h, solution, "exports", format="markdown")
    monkeypatch.setattr(exports, "render", lambda *args: (_ for _ in ()).throw(ValueError("boom")))
    db.rollback()
    deliveries.run_job(failed.json()["run"]["id"])
    assert (
        client.get(
            f"/api/v1/exports/{failed.json()['run']['id']}/download", headers=h()
        ).status_code
        == 404
    )

    objects = {}
    late, _ = start(client, h, solution, "exports", format="markdown")
    monkeypatch.setattr(exports, "render", lambda *args: b"late")

    def put_then_cancel(key, data, mime):
        objects[key] = data
        assert (
            client.post(f"/api/v1/runs/{late.json()['run']['id']}/cancel", headers=h()).status_code
            == 200
        )

    monkeypatch.setattr(storage, "put_original", put_then_cancel)
    monkeypatch.setattr(storage, "delete_original", lambda key: objects.pop(key, None))
    db.rollback()
    deliveries.run_job(late.json()["run"]["id"])
    db.expire_all()
    run = db.get(GenerationRun, UUID(late.json()["run"]["id"]))
    assert run.status == "cancelled" and objects == {}
    assert client.get(f"/api/v1/exports/{run.id}/download", headers=h()).status_code == 404
