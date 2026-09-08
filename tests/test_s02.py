"""S02 real PostgreSQL transactions; fake object store for deterministic failure injection."""

from datetime import UTC, datetime, timedelta
from functools import lru_cache
from io import BytesIO
from uuid import UUID

import pymupdf
import pytest
from docx import Document as WordDocument
from solution_copilot.application import document_jobs, parsing
from solution_copilot.application.errors import AppError
from solution_copilot.domain.models import CustomerAccess, DocumentChunk, Job
from solution_copilot.infrastructure import documents as storage
from sqlalchemy import delete, select
from test_s01 import db as db
from test_s01 import setup as setup


@pytest.fixture(autouse=True)
def objects(monkeypatch):
    values = {}
    monkeypatch.setattr(storage, "put_original", lambda key, data, mime: values.update({key: data}))
    monkeypatch.setattr(storage, "read_original", lambda key: values[key])
    monkeypatch.setattr(storage, "delete_original", lambda key: values.pop(key, None))
    return values


@lru_cache
def fixture_file(extension):
    if extension == ".pdf":
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((40, 40), "Requirements page one")
            pdf.new_page().insert_text((40, 40), "Acceptance page two")
            return pdf.tobytes()
    if extension == ".docx":
        doc = WordDocument()
        doc.add_heading("需求", 1)
        doc.add_paragraph("客户要求单点登录。")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "功能", "验收"
        table.cell(1, 0).text, table.cell(1, 1).text = "登录", "通过"
        output = BytesIO()
        doc.save(output)
        return output.getvalue()
    return "# 需求\n\n客户要求单点登录。\n\n## 验收\n\n登录成功。".encode()


def upload(client, headers, extension=".md", scope="organization", **fields):
    response = client.post(
        "/api/v1/documents",
        headers=headers,
        data={"scope": scope, **fields},
        files={
            "file": ("requirements" + extension, fixture_file(extension), parsing.MIMES[extension])
        },
    )
    assert response.status_code == 202, response.text
    return response.json()


@pytest.mark.parametrize("extension", [".txt", ".md", ".pdf", ".docx"])
def test_four_formats_and_idempotency(setup, db, extension):
    client, h, *_ = setup
    result = upload(client, h(), extension)
    doc_id, job_id = result["document"]["id"], result["job"]["id"]
    duplicate = upload(client, h(), extension)
    assert duplicate["duplicate"] and duplicate["job"]["id"] == job_id
    db.rollback()
    document_jobs.run_job(job_id)
    document_jobs.run_job(job_id)
    response = client.get(f"/api/v1/documents/{doc_id}", headers=h()).json()
    assert response["status"] == "parsed", response
    values = client.get(f"/api/v1/documents/{doc_id}/chunks", headers=h()).json()["items"]
    assert len({c["id"] for c in values}) == len(values) > 0
    assert all(c["generation"] == 1 for c in values)
    if extension == ".pdf":
        assert {c["page_number"] for c in values} == {1, 2}
        assert values[0]["metadata"]["bbox"]
    elif extension == ".docx":
        assert any(c["section_path"] == ["需求"] for c in values)
        assert any(c["metadata"]["table"] and "功能" in c["content"] for c in values)
    else:
        assert any("单点登录" in c["content"] for c in values)
        assert values[0]["metadata"]["line_start"] == 1
    assert client.get(f"/api/v1/documents/{doc_id}/download", headers=h()).content == fixture_file(
        extension
    )


def test_scope_and_execution_revocation(setup, db):
    client, h, customers, users, _ = setup
    response = client.post(
        "/api/v1/documents",
        headers=h("a-member"),
        data={"scope": "organization"},
        files={"file": ("test.txt", b"data", "text/plain")},
    )
    assert response.status_code == 403
    result = upload(client, h("a-member"), scope="customer", customer_id=str(customers[0].id))
    doc_id, job_id = result["document"]["id"], result["job"]["id"]
    for route in [
        f"documents/{doc_id}",
        f"documents/{doc_id}/download",
        f"documents/{doc_id}/chunks",
        f"jobs/{job_id}",
    ]:
        assert client.get("/api/v1/" + route, headers=h("b-owner", 1)).status_code == 404
    assert (
        client.post(f"/api/v1/documents/{doc_id}/disable", headers=h("a-viewer")).status_code == 403
    )
    db.execute(delete(CustomerAccess).where(CustomerAccess.user_id == users["a-member"][0].id))
    db.commit()
    document_jobs.run_job(job_id)
    result = client.get(f"/api/v1/jobs/{job_id}", headers=h()).json()
    assert result["status"] == "failed" and result["error_code"] == "ACCESS_REVOKED"
    assert (
        client.get(f"/api/v1/documents/{doc_id}/download", headers=h("a-member")).status_code == 404
    )
    assert not list(db.scalars(select(DocumentChunk)))


def test_failure_retry_atomic_replace_cancel_delete(setup, db, monkeypatch, objects):
    client, h, *_ = setup
    result = upload(client, h())
    doc_id, job_id = result["document"]["id"], result["job"]["id"]
    db.rollback()
    document_jobs.run_job(job_id)
    old_ids = [c.id for c in db.scalars(select(DocumentChunk))]
    result = client.post(f"/api/v1/documents/{doc_id}/reindex", headers=h()).json()
    real_chunks = document_jobs.chunks

    def broken(*args):
        raise AppError(422, "NO_TEXT", "请补充文本后重试。")

    monkeypatch.setattr(document_jobs, "chunks", broken)
    db.rollback()
    document_jobs.run_job(result["id"])
    assert [c.id for c in db.scalars(select(DocumentChunk))] == old_ids
    retry = client.post(f"/api/v1/jobs/{result['id']}/retry", headers=h()).json()
    monkeypatch.setattr(document_jobs, "chunks", real_chunks)
    db.rollback()
    document_jobs.run_job(retry["id"])
    assert (
        client.post(f"/api/v1/jobs/{result['id']}/retry", headers=h()).json()["id"] == retry["id"]
    )
    assert all(c.generation == 3 for c in db.scalars(select(DocumentChunk)))
    result = client.post(f"/api/v1/documents/{doc_id}/reindex", headers=h()).json()
    assert (
        client.post(f"/api/v1/jobs/{result['id']}/cancel", headers=h()).json()["status"]
        == "cancelled"
    )
    db.rollback()
    document_jobs.run_job(result["id"])
    assert client.get(f"/api/v1/documents/{doc_id}", headers=h()).json()["generation"] == 3
    assert client.delete(f"/api/v1/documents/{doc_id}", headers=h()).status_code == 200
    assert client.get(f"/api/v1/documents/{doc_id}/download", headers=h()).status_code == 404
    db.rollback()
    document_jobs.dispatch_once(lambda _: None)
    assert not objects
    assert not list(db.scalars(select(DocumentChunk)))
    revived = upload(client, h())
    assert revived["document"]["id"] == doc_id and objects


def test_outbox_and_worker_lease_recovery(setup, db):
    client, h, *_ = setup
    result = upload(client, h())
    job_id = UUID(result["job"]["id"])

    def broker_down(_):
        raise ConnectionError()

    db.rollback()
    document_jobs.dispatch_once(broker_down)
    job = db.get(Job, job_id, populate_existing=True)
    assert job.status == "queued"
    job.dispatch_after = datetime.now(UTC) - timedelta(seconds=1)
    job.status, job.attempts = "running", 1
    job.lease_until = datetime.now(UTC) - timedelta(seconds=1)
    db.commit()
    sent = []
    document_jobs.dispatch_once(sent.append)
    assert sent == [str(job_id)]
    document_jobs.run_job(str(job_id))
    db.expire_all()
    assert db.get(Job, job_id).status == "succeeded"
    assert db.get(Job, job_id).attempts == 2


def test_upload_validation_and_structured_chunks():
    for name, mime, data in [
        ("x.exe", "text/plain", b"test"),
        ("x.pdf", "application/pdf", b"not pdf"),
        ("x.docx", parsing.MIMES[".docx"], b"not zip"),
        ("x.txt", "text/plain", b"a\x00b"),
        ("x.txt", "application/pdf", b"test"),
    ]:
        with pytest.raises(AppError):
            parsing.validate_file(name, mime, data)
    long = parsing.chunks(("中" * 1500).encode(), ".txt")
    assert [len(c["content"]) for c in long] == [700, 700, 300]
    table = "|功能|验收|\n|---|---|\n" + "|登录|通过|\n" * 200
    pieces = parsing.chunks(table.encode(), ".md")
    assert len(pieces) > 1 and all(c["content"].startswith("|功能|验收|") for c in pieces)
    with pymupdf.open() as pdf:
        pdf.new_page()
        with pytest.raises(AppError) as err:
            parsing.chunks(pdf.tobytes(), ".pdf")
        assert err.value.code == "NO_TEXT"


def test_cancel_during_parse_fences_publication(setup, db, monkeypatch):
    client, h, *_ = setup
    result = upload(client, h())
    job_id = result["job"]["id"]
    original = document_jobs.chunks

    def cancel_while_parsing(data, extension, checkpoint):
        assert client.post(f"/api/v1/jobs/{job_id}/cancel", headers=h()).status_code == 200
        db.rollback()
        return original(data, extension, checkpoint)

    monkeypatch.setattr(document_jobs, "chunks", cancel_while_parsing)
    db.rollback()
    document_jobs.run_job(job_id)
    assert client.get(f"/api/v1/jobs/{job_id}", headers=h()).json()["status"] == "cancelled"
    assert not list(db.scalars(select(DocumentChunk)))


def test_upload_body_limits_and_invalid_scopes(setup, monkeypatch):
    client, h, customers, *_ = setup
    from solution_copilot.config import get_settings

    settings = get_settings().model_copy(update={"upload_max_bytes": 100})
    monkeypatch.setattr("apps.api.documents.get_settings", lambda: settings)
    monkeypatch.setattr("apps.api.upload_limit.get_settings", lambda: settings)
    monkeypatch.setattr(parsing, "get_settings", lambda: settings)

    def post(scope="organization", data=b"test", **fields):
        return client.post(
            "/api/v1/documents",
            headers=h(),
            data={"scope": scope, **fields},
            files={"file": ("test.txt", data, "text/plain")},
        )

    assert post(data=b"x" * 101).status_code == 413
    assert post(data=b"x" * (1024 * 1024 + 101)).status_code == 413
    assert post("public").status_code == 422
    assert post(customer_id=str(customers[0].id)).status_code == 422
    assert post("customer", customer_id=str(customers[2].id)).status_code == 404


def test_storage_failure_is_durable_and_bounded(setup, db, monkeypatch):
    client, h, *_ = setup

    def broken(*args):
        raise ConnectionError()

    monkeypatch.setattr(storage, "put_original", broken)
    result = upload(client, h())
    job_id = result["job"]["id"]
    db.rollback()
    for _ in range(3):
        document_jobs.run_job(job_id)
    state = client.get(f"/api/v1/jobs/{job_id}", headers=h()).json()
    assert state["status"] == "failed" and state["attempts"] == 3
    assert state["error_code"] == "PROCESSING_FAILED"
