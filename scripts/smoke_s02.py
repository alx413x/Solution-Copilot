"""Live Web -> API -> MinIO -> dispatcher -> Celery -> PG acceptance, own fixtures only."""

import json
import time
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import httpx
import pymupdf
from docx import Document


def main():
    credential = json.loads(Path(".local/demo-credentials.json").read_text())["a-owner"]
    origin = "http://127.0.0.1:3000"
    with httpx.Client(
        base_url=origin, headers={"Origin": origin}, trust_env=False, timeout=60
    ) as web:
        response = web.post("/api/session", json={"token": credential["token"]})
        assert response.status_code == 200, response.text
        marker = "S02 acceptance " + uuid4().hex
        word = Document()
        word.add_heading("验收资料", 1)
        word.add_paragraph(marker)
        stream = BytesIO()
        word.save(stream)
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((40, 40), marker)
            pdf_bytes = pdf.tobytes()
        files = [
            ("txt", marker.encode(), "text/plain"),
            ("md", ("# 验收\n\n" + marker).encode(), "text/markdown"),
            ("pdf", pdf_bytes, "application/pdf"),
            (
                "docx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        ]
        ids = []
        try:
            for extension, data, mime in files:
                payload = {"file": (f"验收资料.{extension}", data, mime)}
                response = web.post(
                    "/api/backend/documents", data={"scope": "organization"}, files=payload
                )
                assert response.status_code == 202, response.text
                result = response.json()
                doc_id, job_id = result["document"]["id"], result["job"]["id"]
                ids.append(doc_id)
                duplicate = web.post(
                    "/api/backend/documents", data={"scope": "organization"}, files=payload
                )
                assert duplicate.json()["duplicate"]
                deadline = time.monotonic() + 60
                while time.monotonic() < deadline:
                    job = web.get(f"/api/backend/jobs/{job_id}").json()
                    if job["status"] in {"failed", "succeeded"}:
                        break
                    time.sleep(0.5)
                assert job["status"] == "succeeded", job
                chunks = web.get(f"/api/backend/documents/{doc_id}/chunks").json()["items"]
                assert marker in "\n".join(c["content"] for c in chunks)
                original = web.get(f"/api/backend/documents/{doc_id}/download")
                assert original.content == data
                assert "no-store" in original.headers["cache-control"]
                assert "attachment" in original.headers["content-disposition"]
                print(f"PASS {extension}: private upload, dedup, worker, locations, download")
            blocked = web.post(
                "/api/backend/documents",
                headers={"Origin": "https://evil.invalid"},
                data={"scope": "organization"},
                files=payload,
            )
            assert blocked.status_code == 403
        finally:
            for doc_id in ids:
                assert web.delete(f"/api/backend/documents/{doc_id}").status_code == 200
                assert web.get(f"/api/backend/documents/{doc_id}/download").status_code == 404
        print("PASS origin guard and immediate deletion visibility")


if __name__ == "__main__":
    main()
