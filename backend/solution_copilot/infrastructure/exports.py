"""Render the saved editor tree, never HTML; references keep chapter-local numbers."""

import re
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from solution_copilot.application.solution_schemas import validate_document

MIMES = {
    "markdown": "text/markdown; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def plain(node):
    if node["type"] == "text":
        return node["text"]
    if node["type"] == "hardBreak":
        return "\n"
    separator = "" if node["type"] in {"paragraph", "heading", "codeBlock"} else "\n"
    return separator.join(plain(c) for c in node.get("content", []))


def missing_references(section):
    used = set(re.findall(r"\[(\d+)\]", plain(section["content"])))
    return used - {c["id"] for c in section["citations"]}


def escape(text):
    return re.sub(r"([\\`*_{}\[\]<>()#+.!|~>-])", r"\\\1", text)


def markdown(node, depth=0):
    kind, children = node["type"], node.get("content", [])
    if kind == "table":
        rows = [
            "| "
            + " | ".join(markdown(cell).strip().replace("\n", "<br>") for cell in row["content"])
            + " |"
            for row in children
        ]
        # Markdown requires a header; preserve every original row when no header was supplied.
        if not all(c["type"] == "tableHeader" for c in children[0]["content"]):
            rows.insert(0, "| " + " | ".join(" " for _ in children[0]["content"]) + " |")
        rows.insert(1, "| " + " | ".join("---" for _ in children[0]["content"]) + " |")
        return "\n".join(rows) + "\n\n"
    if kind == "text":
        value = escape(node["text"])
        for mark in node.get("marks", []):
            if mark["type"] == "code":
                fence = "`" * (
                    max((len(m) for m in re.findall(r"`+", node["text"])), default=0) + 1
                )
                value = fence + " " + node["text"] + " " + fence
            else:
                token = {"bold": "**", "italic": "*", "strike": "~~"}[mark["type"]]
                value = token + value + token
        return value
    if kind == "hardBreak":
        return "  \n"
    if kind == "codeBlock":
        text = plain(node)
        fence = "`" * max(3, max((len(m) + 1 for m in re.findall(r"`+", text)), default=3))
        return (
            fence
            + (node.get("attrs", {}).get("language") or "")
            + "\n"
            + text
            + "\n"
            + fence
            + "\n\n"
        )
    if kind in {"bulletList", "orderedList"}:
        lines = []
        for index, child in enumerate(children, node.get("attrs", {}).get("start", 1)):
            prefix = f"{index}. " if kind == "orderedList" else "- "
            body = markdown(child, depth + 1).strip().splitlines() or [""]
            lines.append(
                prefix + body[0] + "\n" + "\n".join(" " * len(prefix) + s for s in body[1:])
            )
        return "\n".join(lines) + "\n\n"
    body = "".join(markdown(c, depth) for c in children)
    if kind == "heading":
        return "#" * (node["attrs"]["level"] + 1) + " " + body + "\n\n"
    if kind == "paragraph":
        return body + "\n\n"
    if kind == "blockquote":
        return "\n".join("> " + line for line in body.rstrip().splitlines()) + "\n\n"
    if kind == "horizontalRule":
        return "---\n\n"
    return body


def references(sections):
    for index, section in enumerate(sections, 1):
        for c in section["citations"]:
            locator = c.get("locator", {})
            place = " / ".join(str(s) for s in locator.get("section_path", []) or [])
            if locator.get("page_number"):
                place += f" 页 {locator['page_number']}"
            yield (
                f"第 {index} 章 [{c['id']}] {c['title']} · {place}\n"
                f"来源编号：{c['source_id']}\n摘录：{c['quote']}\n"
                "状态：待核实（引用摘录不等于事实已验证）"
            )


def add_inline(paragraph, node):
    if node["type"] == "text":
        run = paragraph.add_run(node["text"])
        marks = {m["type"] for m in node.get("marks", [])}
        run.bold, run.italic = "bold" in marks, "italic" in marks
        run.font.strike = "strike" in marks
        if "code" in marks:
            run.font.name = "Courier New"
    elif node["type"] == "hardBreak":
        paragraph.add_run().add_break()
    else:
        for child in node.get("content", []):
            add_inline(paragraph, child)


def docx_nodes(doc, node, depth=0, prefix=""):
    kind = node["type"]
    if kind == "table":
        table = doc.add_table(rows=0, cols=len(node["content"][0]["content"]))
        table.style = "Table Grid"
        for row in node["content"]:
            cells = table.add_row().cells
            for cell, value in zip(cells, row["content"], strict=True):
                for i, paragraph in enumerate(value["content"]):
                    p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                    add_inline(p, paragraph)
                    if value["type"] == "tableHeader":
                        for run in p.runs:
                            run.bold = True
    elif kind in {"paragraph", "heading", "codeBlock"}:
        style = "Normal"
        if kind == "heading":
            style = f"Heading {node['attrs']['level']}"
        elif kind == "codeBlock":
            style = "No Spacing"
        p = doc.add_paragraph(style=style)
        p.paragraph_format.left_indent = Cm(min(depth, 8) * 0.5)
        if prefix:
            p.add_run(prefix)
        add_inline(p, node)
        if kind == "codeBlock":
            for run in p.runs:
                run.font.name = "Courier New"
                run.font.size = Pt(9)
    elif kind in {"bulletList", "orderedList"}:
        for i, child in enumerate(node["content"], node.get("attrs", {}).get("start", 1)):
            children = child["content"]
            docx_nodes(doc, children[0], depth + 1, f"{i}. " if kind == "orderedList" else "• ")
            for nested in children[1:]:
                docx_nodes(doc, nested, depth + 1)
    elif kind == "horizontalRule":
        doc.add_paragraph("—")
    else:
        for child in node.get("content", []):
            docx_nodes(doc, child, depth + (kind == "blockquote"))


def render(title, version, sections, format):
    for section in sections:
        validate_document(section["content"])
        if missing_references(section):
            raise ValueError("Unknown citation number")
    notice = f"版本 {version} · AI 草稿，所有内容与引用均需人工核实。引用编号按章节独立编号。"
    if format == "markdown":
        parts = ["# " + escape(title), escape(notice)]
        for i, s in enumerate(sections, 1):
            parts += [f"## {i} {escape(s['title'])}", "**待核实**", markdown(s["content"])]
            parts.extend(escape(w) for w in s["warnings"])
        parts += ["## 参考资料", *(escape(r) for r in references(sections))]
        data = ("\n\n".join(parts) + "\n").encode("utf-8")
        data.decode("utf-8")
        return data
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.3)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "No Spacing"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(8)
    section.header.paragraphs[0].text = title
    footer = section.footer.paragraphs[0]
    footer.text = f"版本 {version} · 待人工核实 · "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.add_paragraph(title, "Title")
    doc.add_paragraph(notice)
    for i, s in enumerate(sections, 1):
        doc.add_heading(f"{i} {s['title']}", 1)
        doc.add_paragraph("待核实").runs[0].bold = True
        docx_nodes(doc, s["content"])
        for warning in s["warnings"]:
            doc.add_paragraph(warning)
    doc.add_heading("参考资料", 1)
    for reference in references(sections):
        doc.add_paragraph(reference)
    buffer = BytesIO()
    doc.save(buffer)
    data = buffer.getvalue()
    with ZipFile(BytesIO(data)) as archive:
        if archive.testzip() is not None:
            raise ValueError("Invalid DOCX archive")
    opened = Document(BytesIO(data))
    if opened.paragraphs[0].text != title:
        raise ValueError("DOCX verification failed")
    return data
